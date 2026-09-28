"""Fixed-width text implementation of :class:`DocumentProcessingPort`.

This is the pre-existing extraction and chunking behaviour -- ``python-docx``
for ``.docx``, UTF-8 decoding for the text formats, and fixed-width paragraph
and sentence chunking -- presented through the same port as the structural
parser.

It exists for two reasons:

* it is the pipeline the application used before the structural parser, and
  remains the active one while that parser is held behind ``DOCLING_ENABLED``;
* it is the fallback for formats the structural parser cannot handle, so a
  supported file type is never silently lost when the parser is enabled.

The implementation lives here, behind the port, so the use-case layer only
builds canonical vector records from neutral processing results.
"""

import logging
import re
from io import BytesIO
from pathlib import Path
from typing import Sequence

from docx import Document

from app.domain.ports.document_processing import (
    DocumentProcessingPort,
    ProcessedChunk,
    ProcessedDocument,
)
logger = logging.getLogger("knowlayer.ingestion")

SUPPORTED_EXTENSIONS = {".docx", ".txt", ".md", ".rst"}


def _clean_text(text: str) -> str:
    text = text.replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    text = re.sub(r"\s{2,}", " ", text)
    return text.strip()


def chunk_text(text: str, chunk_size: int = 500) -> list[str]:
    cleaned = _clean_text(text)
    if not cleaned:
        return []

    paragraphs = [part.strip() for part in re.split(r"\n\s*\n+", cleaned) if part.strip()]
    if not paragraphs:
        paragraphs = [cleaned]

    chunks: list[str] = []
    current = ""

    for paragraph in paragraphs:
        if len(paragraph) <= chunk_size:
            if not current:
                current = paragraph
            elif len(current) + 1 + len(paragraph) <= chunk_size:
                current = f"{current} {paragraph}"
            else:
                chunks.append(current.strip())
                current = paragraph
            continue

        sentences = re.split(r"(?<=[.!?])\s+", paragraph)
        for sentence in sentences:
            sentence = sentence.strip()
            if not sentence:
                continue
            if not current:
                current = sentence
            elif len(current) + 1 + len(sentence) <= chunk_size:
                current = f"{current} {sentence}"
            else:
                chunks.append(current.strip())
                current = sentence

    if current.strip():
        chunks.append(current.strip())

    return [chunk for chunk in chunks if chunk.strip()] or [cleaned[:chunk_size]]


def extract_text_from_bytes(content: bytes, file_name: str) -> str:
    suffix = Path(file_name).suffix.lower()

    if suffix == ".docx":
        document = Document(BytesIO(content))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        return "\n\n".join(paragraphs)

    if suffix in {".txt", ".md", ".rst"}:
        return content.decode("utf-8", errors="ignore")

    raise ValueError(f"Unsupported document type: {suffix}")


def extract_text_from_document(file_path: str) -> str:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(file_path)
    return extract_text_from_bytes(path.read_bytes(), path.name)


class PlainTextDocumentProcessor:
    """Parse and chunk documents with the fixed-width text pipeline.

    Reports no document structure: the fixed-width splitter has no notion of
    headings or pages, so ``page_number`` is left unset, the heading hierarchy
    is empty, and the content type is plain text. Those omissions are
    meaningful and are not backfilled.
    """

    def supported_extensions(self) -> Sequence[str]:
        """Extensions the text pipeline has always accepted."""
        return tuple(sorted(SUPPORTED_EXTENSIONS))

    def supports(self, file_name: str) -> bool:
        return Path(file_name).suffix.lower() in SUPPORTED_EXTENSIONS

    def process(self, content: bytes, file_name: str) -> ProcessedDocument:
        """Extract text and split it into fixed-width chunks.

        Raises ``ValueError`` for an unsupported extension, matching the
        convention the ingestion layer maps to a rejected upload.
        """
        if not self.supports(file_name):
            supported = ", ".join(self.supported_extensions())
            suffix = Path(file_name).suffix.lower()
            raise ValueError(
                f"Unsupported document type: {suffix or '(none)'} (supported: {supported})"
            )

        text = extract_text_from_bytes(content, file_name)
        chunks = tuple(
            ProcessedChunk(text=chunk, index=index)
            for index, chunk in enumerate(chunk_text(text))
            if chunk.strip()
        )
        logger.info(
            "plaintext processed %r into %d chunks", file_name, len(chunks)
        )
        return ProcessedDocument(
            chunks=chunks,
            file_type=Path(file_name).suffix.lower().lstrip(".") or None,
        )
