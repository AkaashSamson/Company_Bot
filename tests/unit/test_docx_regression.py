from io import BytesIO
import pytest
from docx import Document
from app.use_cases.ingestion.document_processing import extract_blocks_from_docx

def test_docx_table_duplication_regression():
    """
    Regression test to ensure that a single table in a DOCX document
    is not duplicated into multiple DocumentBlocks during extraction.
    """
    # Create a real DOCX file in memory
    doc = Document()
    doc.add_paragraph("This is a paragraph before the table.")
    
    table = doc.add_table(rows=3, cols=2)
    
    # Header row
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Header 1"
    hdr_cells[1].text = "Header 2"
    
    # Data row 1
    row_cells = table.rows[1].cells
    row_cells[0].text = "Data 1"
    row_cells[1].text = "Data 2"
    
    # Data row 2
    row_cells = table.rows[2].cells
    row_cells[0].text = "Data 3"
    row_cells[1].text = "Data 4"
    
    doc.add_paragraph("This is a paragraph after the table.")
    
    # Save to bytes
    doc_bytes = BytesIO()
    doc.save(doc_bytes)
    doc_bytes.seek(0)
    
    # Extract blocks
    blocks = extract_blocks_from_docx(doc_bytes.read())
    
    # We expect 3 blocks: Paragraph, Table, Paragraph.
    # If the duplication bug exists, it would produce something like:
    # Paragraph, Table (1 row), Table (2 rows), Paragraph
    
    table_blocks = [b for b in blocks if b.block_type == 'table']
    
    # Verify exactly one table block is produced
    assert len(table_blocks) == 1, f"Expected 1 table block, but got {len(table_blocks)}"
    
    # Verify the table contains the correct number of mapped data rows (excluding the header)
    assert len(table_blocks[0].table_rows) == 2
    assert "Header 1: Data 1" in table_blocks[0].table_rows[0]
